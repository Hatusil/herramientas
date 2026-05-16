"""
Processor: Funciones para limpiar metadatos de archivos.
"""
import os
import logging
from pathlib import Path
from typing import Dict, Any

# Importar funciones compartidas de core (máxima #4: Lo que me gusta en una herramienta debe estar en las demás)
from core.utils import get_output_path, validate_input_file, validate_file_size

# Métricas
from core.metrics import Counter, Timer, increment

logger = logging.getLogger(__name__)

MAX_SCRUB_SIZE_MB = 2000  # 2GB max for scrubbing

# Contadores de operaciones
scrubber_operations_total = Counter('scrubber_operations_total')
scrubber_errors = Counter('scrubber_errors')

# =============================================================================
# IMÁGENES - EXIF
# =============================================================================

def get_image_metadata(file_path: str) -> Dict[str, Any]:
    """
    Obtiene metadatos EXIF de una imagen.
    
    Args:
        file_path: Ruta al archivo de imagen
        
    Returns:
        dict: Metadatos encontrados
    """
    if not os.path.exists(file_path):
        return {'success': False, 'error': 'Archivo no encontrado'}
    
    ext = Path(file_path).suffix.lower()
    if ext not in ['.jpg', '.jpeg', '.png', '.tiff', '.webp']:
        return {'success': False, 'error': 'Formato no soportado para imágenes'}
    
    try:
        from PIL import Image
        
        with Image.open(file_path) as img:
            info = {}
            
            # EXIF data
            exif_data = img.getexif()
            if exif_data:
                for tag_id, value in exif_data.items():
                    try:
                        tag_name = str(tag_id)
                        info[f'EXIF_{tag_name}'] = str(value)[:100]
                    except Exception as e:
                        logger.warning(f"Error reading EXIF tag {tag_id}: {e}")
            
            # Info básica
            info['format'] = img.format
            info['size'] = img.size
            info['mode'] = img.mode
            
            # Verificar si tiene EXIF
            has_exif = len(exif_data) > 0 if exif_data else False
            
            # Tratar de obtener datos específicos
            try:
                exif = img.getexif()
                if exif:
                    # Fecha
                    if 0x9003 in exif:  # DateTimeOriginal
                        info['date_taken'] = str(exif[0x9003])
                    # Cámara
                    if 0x010F in exif:  # Make
                        info['camera_make'] = str(exif[0x010F])
                    if 0x0110 in exif:  # Model
                        info['camera_model'] = str(exif[0x0110])
                    # GPS
                    if 0x8825 in exif:  # GPSInfo
                        info['has_gps'] = True
            except Exception as e:
                logger.warning(f"Error getting EXIF metadata: {e}")
            
            return {
                'success': True,
                'file_name': os.path.basename(file_path),
                'has_exif': has_exif,
                'metadata': info
            }
            
    except Exception as e:
        logger.error(f"Error leyendo metadatos de imagen: {e}")
        return {'success': False, 'error': str(e)}


def clean_image_metadata(file_path: str, options: Dict[str, Any] = None) -> Dict[str, Any]:
    """
    Limpia metadatos de una imagen.
    
    Args:
        file_path: Ruta al archivo de imagen
        options: Opciones {'remove_all': True, 'remove_gps': True, 'remove_date': True}
        
    Returns:
        dict: Resultado de la operación
    """
    with Timer('scrubber.clean_image_metadata'):
        if options is None:
            options = {}
        
        if not os.path.exists(file_path):
            increment('scrubber_errors')
            return {'success': False, 'error': 'Archivo no encontrado', 'output_files': []}
        
        ext = Path(file_path).suffix.lower()
        if ext not in ['.jpg', '.jpeg']:
            increment('scrubber_errors')
            return {'success': False, 'error': 'Solo se soporta JPG/JPEG para eliminar EXIF', 'output_files': []}
        
        try:
            # Intentar usar piexif
            try:
                import piexif
                
                # Cargar EXIF actual
                try:
                    exif_dict = piexif.load(file_path)
                except Exception as e:
                    logger.warning(f"Could not load existing EXIF: {e}")
                    exif_dict = {"0th": {}, "Exif": {}, "GPS": {}, "1st": {}, "thumbnail": None}
                
                remove_all = options.get('remove_all', True)
                
                if remove_all:
                    # Eliminar todo
                    exif_dict = {"0th": {}, "Exif": {}, "GPS": {}, "1st": {}, "thumbnail": None}
                else:
                    # Opciones específicas
                    if options.get('remove_gps', False):
                        exif_dict['GPS'] = {}
                    if options.get('remove_date', False):
                        if '0th' in exif_dict and piexif.ImageIFD.DateTime in exif_dict['0th']:
                            del exif_dict['0th'][piexif.ImageIFD.DateTime]
                
                # Guardar
                exif_bytes = piexif.dump(exif_dict)
                
                # Leer imagen y guardar sin EXIF o con EXIF limpiado
                from PIL import Image
                with Image.open(file_path) as img:
                    output_path = get_output_path(file_path, '_clean', _exists_ok=False)
                    img.save(output_path, "jpeg", exif=exif_bytes, quality=95)
                
                increment('scrubber_operations_total')
                
                return {
                    'success': True,
                    'message': 'Metadatos eliminados',
                    'output_files': [output_path],
                    'error': None
                }
                
            except ImportError:
                # Fallback: guardar sin EXIF
                from PIL import Image
                with Image.open(file_path) as img:
                    output_path = get_output_path(file_path, '_clean', _exists_ok=False)
                    data = list(img.getdata())
                    img_no_exif = Image.new(img.mode, img.size)
                    img_no_exif.putdata(data)
                    img_no_exif.save(output_path, "JPEG", quality=95)
                
                increment('scrubber_operations_total')
                
                return {
                    'success': True,
                    'message': 'EXIF eliminado (sin piexif)',
                    'output_files': [output_path],
                    'error': None
                }
                
        except Exception as e:
            increment('scrubber_errors')
            logger.error(f"Error limpiando metadatos: {e}")
            return {'success': False, 'error': str(e), 'output_files': []}


# =============================================================================
# DOCUMENTOS - DOCX
# =============================================================================

def get_docx_metadata(file_path: str) -> Dict[str, Any]:
    """
    Obtiene metadatos de un archivo DOCX.
    """
    if not os.path.exists(file_path):
        return {'success': False, 'error': 'Archivo no encontrado'}
    
    if not file_path.lower().endswith('.docx'):
        return {'success': False, 'error': 'No es un archivo DOCX'}
    
    try:
        from docx import Document
        
        # Abrir el documento
        doc = Document(file_path)
        
        # Obtener propiedades core
        core_props = doc.core_properties
        
        info = {
            'title': core_props.title or '',
            'author': core_props.author or '',
            'subject': core_props.subject or '',
            'keywords': core_props.keywords or '',
            'created': str(core_props.created) if core_props.created else '',
            'modified': str(core_props.modified) if core_props.modified else '',
            'last_modified_by': core_props.last_modified_by or '',
            'revision': core_props.revision or 0,
        }
        
        # Verificar si hay metadatos
        has_metadata = any(v for k, v in info.items() if k != 'revision')
        
        return {
            'success': True,
            'file_name': os.path.basename(file_path),
            'has_metadata': has_metadata,
            'metadata': info
        }
        
    except Exception as e:
        logger.error(f"Error leyendo metadatos DOCX: {e}")
        return {'success': False, 'error': str(e)}


def clean_docx(file_path: str) -> Dict[str, Any]:
    """
    Limpia metadatos de un archivo DOCX.
    """
    with Timer('scrubber.clean_docx'):
        if not os.path.exists(file_path):
            increment('scrubber_errors')
            return {'success': False, 'error': 'Archivo no encontrado', 'output_files': []}
        
        if not file_path.lower().endswith('.docx'):
            increment('scrubber_errors')
            return {'success': False, 'error': 'No es un archivo DOCX', 'output_files': []}
        
        try:
            from docx import Document
            
            doc = Document(file_path)
            
            # Limpiar propiedades core
            core_props = doc.core_properties
            core_props.title = ''
            core_props.author = ''
            core_props.subject = ''
            core_props.keywords = ''
            core_props.last_modified_by = ''
            core_props.comments = ''
            
            # Guardar
            output_path = get_output_path(file_path, '_clean', _exists_ok=False)
            doc.save(output_path)
            
            increment('scrubber_operations_total')
            
            return {
                'success': True,
                'message': 'Metadatos DOCX eliminados',
                'output_files': [output_path],
                'error': None
            }
            
        except Exception as e:
            increment('scrubber_errors')
            logger.error(f"Error limpiando DOCX: {e}")
            return {'success': False, 'error': str(e), 'output_files': []}


# =============================================================================
# DOCUMENTOS - XLSX
# =============================================================================

def get_xlsx_metadata(file_path: str) -> Dict[str, Any]:
    """
    Obtiene metadatos de un archivo XLSX.
    """
    if not os.path.exists(file_path):
        return {'success': False, 'error': 'Archivo no encontrado'}
    
    if not file_path.lower().endswith('.xlsx'):
        return {'success': False, 'error': 'No es un archivo XLSX'}
    
    try:
        import openpyxl
        
        wb = openpyxl.load_workbook(file_path)
        
        props = wb.properties
        
        info = {
            'title': props.title or '',
            'author': props.creator or '',
            'subject': props.subject or '',
            'keywords': props.keywords or '',
            'created': str(props.created) if props.created else '',
            'modified': str(props.modified) if props.modified else '',
            'lastModifiedBy': props.lastModifiedBy or '',
        }
        
        has_metadata = any(v for v in info.values())
        
        return {
            'success': True,
            'file_name': os.path.basename(file_path),
            'has_metadata': has_metadata,
            'metadata': info
        }
        
    except Exception as e:
        logger.error(f"Error leyendo metadatos XLSX: {e}")
        return {'success': False, 'error': str(e)}


def clean_xlsx(file_path: str) -> Dict[str, Any]:
    """
    Limpia metadatos de un archivo XLSX.
    """
    with Timer('scrubber.clean_xlsx'):
        if not os.path.exists(file_path):
            increment('scrubber_errors')
            return {'success': False, 'error': 'Archivo no encontrado', 'output_files': []}
        
        if not file_path.lower().endswith('.xlsx'):
            increment('scrubber_errors')
            return {'success': False, 'error': 'No es un archivo XLSX', 'output_files': []}
        
        try:
            import openpyxl
            
            wb = openpyxl.load_workbook(file_path)
            
            # Limpiar propiedades
            props = wb.properties
            props.title = ''
            props.creator = ''
            props.subject = ''
            props.keywords = ''
            props.lastModifiedBy = ''
            
            # Guardar
            output_path = get_output_path(file_path, '_clean', _exists_ok=False)
            wb.save(output_path)
            
            increment('scrubber_operations_total')
            
            return {
                'success': True,
                'message': 'Metadatos XLSX eliminados',
                'output_files': [output_path],
                'error': None
            }
            
        except Exception as e:
            increment('scrubber_errors')
            logger.error(f"Error limpiando XLSX: {e}")
            return {'success': False, 'error': str(e), 'output_files': []}


# get_output_path() importado desde core.utils (máxima #4)